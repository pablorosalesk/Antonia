"""Helpers comunes para generar los documentos Word con python-docx."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
BASE = 10  # mismo tamaño que el informe original (Calibri 10)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x7A, 0x1F, 0x3D)  # granate sobrio para títulos

# Colores de esfera (tonos del informe original)
C_GENERAL = "F4CCCC"   # rosado
C_NUTRI = "D9D2E9"     # lila
C_COGN = "CFE2F3"      # celeste
C_FUNC = "D9EAD3"      # verde
C_SOCIAL = "FCE5CD"    # durazno
C_HEAD = "EFEFEF"


def _shade(el, color):
    pPr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    pPr.append(shd)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)


def add_runs(p, text, size=None, color=None, bold=None, italic=None):
    """Texto con **negrita**, _cursiva_ no soportada; [[...]] = marcador amarillo (pendiente)."""
    parts = re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text)
    for part in parts:
        if not part:
            continue
        hl = False
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]; b = True
        elif part.startswith("[[") and part.endswith("]]"):
            part = part[2:-2]; hl = True
        r = p.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(size if (size and size >= 12) else BASE)
        if color is not None: r.font.color.rgb = color
        if b: r.bold = True
        if italic: r.italic = True
        if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


class Doc:
    def __init__(self, header_text):
        self.d = Document()
        self.fig = 0
        self.tab = 0
        st = self.d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(10); st.font.color.rgb = INK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(4); pf.space_before = Pt(0); pf.line_spacing = 1.12
        for name, size in (("Heading 1", 10), ("Heading 2", 10), ("Heading 3", 10)):
            h = self.d.styles[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = INK; h.font.italic = False
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), FONT)
            h.paragraph_format.keep_with_next = True
        s = self.d.sections[0]
        s.page_width, s.page_height = Emu(12240 * 635), Emu(15840 * 635)  # carta
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.8)
        s.header_distance = Cm(0.9); s.footer_distance = Cm(0.8)
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]
        add_runs(hp, header_text, size=8, color=MUTED)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(fp, "Página ", size=8, color=MUTED); self._field(fp, "PAGE")
        add_runs(fp, " de ", size=8, color=MUTED); self._field(fp, "NUMPAGES")

    @property
    def width(self):
        s = self.d.sections[0]
        return s.page_width - s.left_margin - s.right_margin

    def _field(self, p, code):
        r = p.add_run(); r.font.size = Pt(BASE); r.font.color.rgb = MUTED; r.font.name = FONT
        for t, txt in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
            elif txt == code:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = f" {code} "; r._r.append(e)
            else:
                e = OxmlElement("w:t"); e.text = txt; r._r.append(e)

    # --- bloques -------------------------------------------------------
    def section(self, title, color):
        p = self.d.add_paragraph(style="Heading 1")
        add_runs(p, title.upper(), size=11.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.space_before = Pt(14); pf.space_after = Pt(8)
        _shade(p._p, color)
        return p

    def h2(self, title, color=None):
        p = self.d.add_paragraph(style="Heading 2")
        add_runs(p, title)
        pf = p.paragraph_format; pf.space_before = Pt(10); pf.space_after = Pt(4)
        if color:
            _shade(p._p, color)
        else:
            self._bottom_border(p)
        return p

    def h3(self, title):
        p = self.d.add_paragraph(style="Heading 3")
        add_runs(p, title, color=ACCENT)
        pf = p.paragraph_format; pf.space_before = Pt(6); pf.space_after = Pt(2)
        return p

    def _bottom_border(self, p):
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "1"); e.set(qn("w:color"), "BFBFBF")
        b.append(e); pPr.append(b)

    def p(self, text, justify=True, size=None, color=None, italic=None, after=None, align=None, keep=False):
        p = self.d.add_paragraph()
        add_runs(p, text, size=size, color=color, italic=italic)
        if align is not None:
            p.alignment = align
        elif justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def label(self, label, text):
        return self.p(f"**{label}:** {text}")

    def bullets(self, items, style="List Bullet"):
        for it in items:
            p = self.d.add_paragraph(style=style)
            add_runs(p, it)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, f"{i}.  " + it)

    def table(self, rows, widths_cm, header=True, head_color=C_HEAD, size=9, bold_first_col=False,
              caption=None, col_colors=None, align_center_cols=()):
        if caption:
            self.tab += 1
            cp = self.p(f"**Tabla {self.tab}.** {caption}", justify=False, size=9, color=MUTED, after=3, keep=True)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(t)
        for ci, w in enumerate(widths_cm):
            t.columns[ci].width = Cm(w)
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if header and ri == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for ci, val in enumerate(row):
                c = tr.cells[ci]; c.width = Cm(widths_cm[ci])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                c.paragraphs[0].text = ""
                lines = val if isinstance(val, list) else [val]
                for li, line in enumerate(lines):
                    para = c.paragraphs[0] if li == 0 else c.add_paragraph()
                    para.paragraph_format.space_after = Pt(1)
                    para.paragraph_format.line_spacing = 1.05
                    b = (header and ri == 0) or (bold_first_col and ci == 0)
                    add_runs(para, str(line), size=size, bold=b or None)
                    if ci in align_center_cols or (header and ri == 0):
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if header and ri == 0:
                    shade_cell(c, head_color)
                elif col_colors and ci in col_colors:
                    shade_cell(c, col_colors[ci])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, path, caption, width_cm):
        self.fig += 1
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.p(f"**Figura {self.fig}.** {caption}", justify=False, size=8.5, color=MUTED,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    def figures_side(self, paths, caption, width_cm):
        self.fig += 1
        t = self.d.add_table(rows=1, cols=len(paths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, pth in zip(t.rows[0].cells, paths):
            para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(pth, width=Cm(width_cm))
        self.p(f"**Figura {self.fig}.** {caption}", justify=False, size=8.5, color=MUTED,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    def result(self, text):
        """Caja destacada con el resultado de una escala."""
        p = self.d.add_paragraph(); add_runs(p, text, size=10)
        _shade(p._p, "F3F3F3")
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:left")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "18"); e.set(qn("w:space"), "6"); e.set(qn("w:color"), "7A1F3D")
        b.append(e); pPr.append(b)
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        return p

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        """Todo justificado (excepto párrafos que solo contienen imágenes)."""
        def fix(pars):
            for p in pars:
                if p._p.xpath(".//w:drawing"):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fix(self.d.paragraphs)
        for t in self.d.tables:
            for row in t.rows:
                for c in row.cells:
                    fix(c.paragraphs)
        s = self.d.sections[0]
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            fix(part.paragraphs)
        self.d.save(path)
